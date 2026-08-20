"""
Shared fixtures: every test app runs against tmp storage dirs and a tmp SQLite
file, injected through the RQA_* env overrides read by get_settings().
"""

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('RQA_DATA_DIR', str(tmp_path / 'data'))
    monkeypatch.setenv('RQA_RESULTS_DIR', str(tmp_path / 'results'))
    monkeypatch.setenv('RQA_REFERENCE_DIR', str(tmp_path / 'reference'))
    monkeypatch.setenv('RQA_DB_URL', f"sqlite:///{tmp_path / 'admin.db'}")
    # Jobs run synchronously inside the request: assertions see final state
    monkeypatch.setenv('RQA_QUEUE_INLINE', '1')
    from src.core.config import get_settings
    get_settings.cache_clear()
    from src.main import create_app
    with TestClient(create_app()) as c:
        yield c
    get_settings.cache_clear()


def build_form_xlsx(path, n_rows=12, step_m=10, road='Т-9999', lane=1, category=2,
                    direction='Прямий', duplicate_ch9_10=True, break_continuity_at=None):
    """Miniature profilometer form with the real sheet layout (header row 22)."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(row=2, column=1, value="Об'єкт (ділянка):"); ws.cell(row=2, column=6, value=road)
    ws.cell(row=6, column=1, value='Технічна категорія'); ws.cell(row=6, column=6, value=category)
    ws.cell(row=6, column=10, value='Номер смуги руху:'); ws.cell(row=6, column=14, value=lane)
    ws.cell(row=7, column=1, value='Напрям руху:'); ws.cell(row=7, column=6, value=direction)
    header = ['км', 'м', 'км', 'м'] + [f'канал {i}' for i in range(1, 11)] + \
             ['широта', 'довгота', 'висота', 'широта', 'довгота', 'висота']
    for col, text in enumerate(header, start=1):
        ws.cell(row=22, column=col, value=text)
    for i in range(n_rows):
        start = i * step_m
        end = start + step_m + (5 if break_continuity_at == i else 0)
        iri = [1.0 + 0.1 * i + 0.01 * ch for ch in range(1, 9)]
        ch8 = iri[-1]
        iri += [ch8, ch8] if duplicate_ch9_10 else [ch8 + 0.5, ch8 + 0.6]
        row = [start // 1000, start % 1000, end // 1000, end % 1000, *iri,
               50.0 + i * 1e-4, 30.0 + i * 1e-4, 120.0,
               50.0 + (i + 1) * 1e-4, 30.0 + (i + 1) * 1e-4, 120.0]
        for col, v in enumerate(row, start=1):
            ws.cell(row=23 + i, column=col, value=v)
    wb.save(path)
    return path

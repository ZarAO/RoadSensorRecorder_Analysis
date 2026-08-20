import math

import pytest
from tests.conftest import build_form_xlsx


def test_parse_happy_path(tmp_path):
    from src.services.reference_forms import parse_form_xlsx
    p = build_form_xlsx(tmp_path / 'f.xlsx', n_rows=12, step_m=10)
    form = parse_form_xlsx(str(p))
    assert form.road_name == 'Т-9999' and form.lane == 1 and form.category == 2
    assert form.step_m == 10 and len(form.intervals) == 12
    assert form.chainage_span_m == 120
    assert list(form.intervals.columns)[:4] == ['km_start', 'm_start', 'km_end', 'm_end']
    assert form.bbox[0] <= 50.0 and form.bbox[2] >= 50.0012
    assert any('канал' in w or 'channel' in w for w in form.warnings)  # ch9/10 duplicate warning


def test_parse_no_header_raises(tmp_path):
    from src.services.reference_forms import parse_form_xlsx
    import openpyxl
    wb = openpyxl.Workbook(); wb.active.cell(row=1, column=1, value='garbage')
    p = tmp_path / 'bad.xlsx'; wb.save(p)
    with pytest.raises(ValueError, match='header'):
        parse_form_xlsx(str(p))


def test_parse_continuity_warning(tmp_path):
    from src.services.reference_forms import parse_form_xlsx
    p = build_form_xlsx(tmp_path / 'gap.xlsx', break_continuity_at=5)
    form = parse_form_xlsx(str(p))
    assert any('gap' in w.lower() or 'розрив' in w.lower() for w in form.warnings)


def test_parse_drops_non_finite_row(tmp_path):
    """float('nan') parses fine — such a row must be dropped, not carried into the stats."""
    from src.services.reference_forms import parse_form_xlsx
    import openpyxl
    p = build_form_xlsx(tmp_path / 'nan.xlsx', n_rows=12, step_m=10)
    wb = openpyxl.load_workbook(p)
    wb.active.cell(row=23 + 3, column=5, value='nan')   # iri_ch1 of the 4th data row
    wb.save(p)

    form = parse_form_xlsx(str(p))
    assert len(form.intervals) == 11
    assert not form.intervals.isna().to_numpy().any()   # no NaN leaked into the table
    assert any('NaN' in w for w in form.warnings)
    assert math.isfinite(form.step_m) and math.isfinite(form.chainage_span_m)
    assert all(math.isfinite(v) for v in form.bbox)


def test_loadable_by_study_matcher(tmp_path):
    """The stored CSV must feed profilometer_validation.match.load_form_10m unchanged."""
    from src.services.reference_forms import parse_form_xlsx
    from profilometer_validation.match import load_form_10m
    p = build_form_xlsx(tmp_path / 'f.xlsx')
    form = parse_form_xlsx(str(p))
    csv = tmp_path / 'intervals_10m.csv'
    form.intervals.to_csv(csv, index=False, encoding='utf-8', lineterminator='\n')
    table = load_form_10m(str(csv))
    assert {'iri_ref_10', 'lat_mid', 'lon_mid', 'chainage_m'} <= set(table.columns)
    assert len(table) == 12

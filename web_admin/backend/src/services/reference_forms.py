"""
Profilometer xlsx form parser: reads the official IRI measurement form
(fixed 20-column layout, metadata rows above a labelled header row) into a
structured reference table. The intervals column set matches the
profilometer_validation study's derived-CSV format exactly, so a stored CSV
feeds `profilometer_validation.match.load_form_10m` unchanged.

Structural failures (unreadable file, missing header, too few data rows)
raise ValueError; everything else is tolerated and surfaced as a warning.
"""

import dataclasses
import os

import numpy as np
import openpyxl
import pandas as pd

# Positional column names for the 20-cell data row, in form order.
DATA_COLUMNS = [
    'km_start', 'm_start', 'km_end', 'm_end',
    *[f'iri_ch{i}' for i in range(1, 11)],
    'lat_start', 'lon_start', 'alt_start',
    'lat_end', 'lon_end', 'alt_end',
]

_HEADER_SEARCH_ROWS = 40
_METADATA_LABELS = {
    'road_name': "Об'єкт",
    'category': 'Технічна категорія',
    'lane': 'Номер смуги',
    'direction': 'Напрям руху',
}


@dataclasses.dataclass
class ParsedForm:
    road_name: str
    direction: str | None
    lane: int | None
    category: int | None
    step_m: float
    intervals: pd.DataFrame
    chainage_span_m: float
    bbox: list
    warnings: list


def _find_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=_HEADER_SEARCH_ROWS):
        cells = [c.value for c in row[:5]]
        first_four = tuple(
            str(v).strip() if v is not None else None for v in cells[:4]
        )
        fifth = str(cells[4]).strip() if len(cells) > 4 and cells[4] is not None else ''
        if first_four == ('км', 'м', 'км', 'м') and fifth.startswith('канал'):
            return row[0].row
    raise ValueError("header row not found: expected 'км|м|км|м|канал 1…'")


def _find_label_value(ws, header_row: int, label_substring: str):
    """First non-empty cell to the right of the first cell containing the label."""
    for row in ws.iter_rows(min_row=1, max_row=header_row - 1):
        for idx, cell in enumerate(row):
            if isinstance(cell.value, str) and label_substring in cell.value:
                for next_cell in row[idx + 1:]:
                    if next_cell.value is not None and str(next_cell.value).strip() != '':
                        return next_cell.value
                return None
    return None


def _coerce_int(value, field_name: str, warnings: list):
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        warnings.append(f"не вдалося розпізнати {field_name}: {value!r}")
        return None


def parse_form_xlsx(path: str) -> ParsedForm:
    """Parse a profilometer xlsx form into a ParsedForm reference table."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
    except Exception as exc:
        raise ValueError(f"не вдалося прочитати xlsx: {exc}") from exc

    header_row = _find_header_row(ws)

    warnings: list = []
    road_name = _find_label_value(ws, header_row, _METADATA_LABELS['road_name'])
    if not road_name:
        road_name = os.path.splitext(os.path.basename(path))[0]
    direction = _find_label_value(ws, header_row, _METADATA_LABELS['direction'])
    lane = _coerce_int(
        _find_label_value(ws, header_row, _METADATA_LABELS['lane']), 'номер смуги руху', warnings
    )
    category = _coerce_int(
        _find_label_value(ws, header_row, _METADATA_LABELS['category']), 'технічну категорію', warnings
    )

    n_cols = len(DATA_COLUMNS)
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1):
        row_idx = row[0].row
        if row[0].value is None:
            break
        raw = [cell.value for cell in row[:n_cols]]
        try:
            rows.append([float(v) for v in raw])
        except (TypeError, ValueError):
            warnings.append(f"рядок {row_idx}: нечислові дані, пропущено")

    if len(rows) < 5:
        raise ValueError(f"too few data rows: found {len(rows)}, need at least 5")

    intervals = pd.DataFrame(rows, columns=DATA_COLUMNS)

    chain_start = intervals['km_start'] * 1000.0 + intervals['m_start']
    chain_end = intervals['km_end'] * 1000.0 + intervals['m_end']

    step_diffs = (chain_end - chain_start).abs()
    step_m = float(np.median(step_diffs))
    off_step = step_diffs[~np.isclose(step_diffs, step_m)]
    if len(off_step) > 0:
        warnings.append(f"{len(off_step)} рядків з кроком, відмінним від медіанного {step_m} м")

    gap_mask = chain_start.iloc[1:].to_numpy() != chain_end.iloc[:-1].to_numpy()
    n_gaps = int(np.sum(gap_mask))
    if n_gaps > 0:
        warnings.append(f'{n_gaps} chainage gap/overlap rows')

    ch8 = intervals['iri_ch8']
    if (intervals['iri_ch9'] == ch8).all() and (intervals['iri_ch10'] == ch8).all():
        warnings.append('канали 9–10 дублюють канал 8; еталон рахується з каналів 1–8')

    chainage_span_m = float(abs(chain_end.iloc[-1] - chain_start.iloc[0]))

    lats = pd.concat([intervals['lat_start'], intervals['lat_end']])
    lons = pd.concat([intervals['lon_start'], intervals['lon_end']])
    bbox = [float(lats.min()), float(lons.min()), float(lats.max()), float(lons.max())]

    return ParsedForm(
        road_name=str(road_name),
        direction=str(direction) if direction is not None else None,
        lane=lane,
        category=category,
        step_m=step_m,
        intervals=intervals,
        chainage_span_m=chainage_span_m,
        bbox=bbox,
        warnings=warnings,
    )
